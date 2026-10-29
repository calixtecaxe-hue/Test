"""Version Excel de l'inventaire de l'audit final (29 octobre 2026).

Lit docs/audit-final/inventaire.json, anomalies.json et chiffres.json (produits par audit-final-rapport.mjs, à lancer d'abord)
et écrit docs/audit-final/CAXE-inventaire-questions.xlsx : un onglet « Inventaire » (une ligne par question, filtres, en-tête
figé), un onglet « Anomalies » et un onglet « Synthèse ». Rien n'est recopié à la main.

    python3 scripts/verification/audit-final-xlsx.py
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RACINE = Path(__file__).resolve().parents[2]
DOSSIER = RACINE / "docs" / "audit-final"
inventaire = json.loads((DOSSIER / "inventaire.json").read_text(encoding="utf-8"))
anomalies = json.loads((DOSSIER / "anomalies.json").read_text(encoding="utf-8"))["anomalies"]
chiffres = json.loads((DOSSIER / "chiffres.json").read_text(encoding="utf-8"))
banque = {q["donnee"]: q for q in json.loads((RACINE / "data" / "banque" / "questions.json").read_text(encoding="utf-8"))}

ENTETE = Font(bold=True, color="FFFFFF")
FOND = PatternFill("solid", fgColor="1E4E96")
HAUT = Alignment(vertical="top", wrap_text=True)
LARGEURS = {"identifiant": 12, "agent": 13, "theme": 22, "niveau": 11, "question": 50, "information": 40, "profils": 28, "conditions": 34, "usage": 50, "usagePrevu": 28, "usageExecute": 28, "implementation": 24, "effort": 8, "verification": 60, "statut": 26, "recommandation": 16, "categories": 22, "anomalies": 12, "propositions": 60, "conseilGenerique": 34}


def feuille(classeur, titre, colonnes, lignes, largeurs):
    f = classeur.create_sheet(titre)
    f.append(colonnes)
    for c in f[1]:
        c.font, c.fill, c.alignment = ENTETE, FOND, HAUT
    for ligne in lignes:
        f.append(ligne)
    for i, largeur in enumerate(largeurs, start=1):
        f.column_dimensions[get_column_letter(i)].width = largeur
    for rang in f.iter_rows(min_row=2):
        for c in rang:
            c.alignment = HAUT
    f.freeze_panes = "B2"
    f.auto_filter.ref = f.dimensions
    return f


classeur = Workbook()
classeur.remove(classeur.active)

cles = [c[0] for c in inventaire["colonnes"]]
feuille(
    classeur,
    "Inventaire",
    [c[1] for c in inventaire["colonnes"]],
    [[("oui" if l[k] else "non") if isinstance(l[k], bool) else l[k] for k in cles] for l in inventaire["lignes"]],
    [LARGEURS.get(k, 20) for k in cles],
)

lignes_anomalies = []
for a in anomalies:
    textes = "\n".join(f"{c} : « {' '.join((banque.get(c, {}).get('texte') or '').split())} »" for c in a["codes"])
    lignes_anomalies.append([a["id"], a["titre"], a["type"], a["correction"]["etat"], textes or "toutes les questions concernées (voir le problème)", a["probleme"], a["consequence"], a["correction"]["texte"], a["controle"]])
feuille(classeur, "Anomalies", ["Identifiant", "Anomalie", "Type", "État", "Questions (texte exact)", "Problème", "Conséquence", "Correction effectuée ou proposée", "Contrôle"], lignes_anomalies, [10, 30, 14, 12, 60, 60, 50, 60, 50])

synthese = [
    ["Questions recensées", chiffres["recensees"]["total"]],
    ["dont banque", chiffres["recensees"]["banque"]],
    ["dont cadrage", chiffres["recensees"]["cadrage"]],
    ["dont en attente", chiffres["recensees"]["enAttente"]],
    ["Évaluées une à une", chiffres["verifiees"]["evaluees"]],
    ["Relues", chiffres["verifiees"]["relues"]],
    ["Assemblées par le vrai code et contrôlées par le balayage", chiffres["verifiees"]["assemblees"]],
    ["Affichées et manipulées dans l'application", chiffres["verifiees"]["manipuleesNavigateur"]],
    ["Jamais posées pendant l'audit (brief de création de contenu)", chiffres["verifiees"]["nonPoseesPendantLAudit"]],
]
synthese += [[f"Statut : {s}", n] for s, n in chiffres["statuts"].items()]
synthese += [["dont exploitation à développer (toutes catégories)", chiffres["aDevelopper"]["total"]]]
synthese += [[f"Implémentation : {i}", sum(v.values())] for i, v in chiffres["implementation"].items()]
synthese += [[f"Anomalies : {etat}", ", ".join(ids)] for etat, ids in chiffres["anomalies"]["parEtat"].items()]
feuille(classeur, "Synthèse", ["", "Nombre"], synthese, [60, 40])

sortie = DOSSIER / "CAXE-inventaire-questions.xlsx"
classeur.save(sortie)
print(f"{sortie.relative_to(RACINE)} : {len(inventaire['lignes'])} questions, {len(anomalies)} anomalies.")

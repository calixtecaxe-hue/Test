#!/usr/bin/env python3
"""Convertit le classeur « CAXE_Themes_Chiffres_4_agents » en fichiers de données
(`data/*.json`). Le classeur reste l'outil d'écriture ; les JSON sont ce que lit
le code. Les relancer ne change rien si le classeur n'a pas changé.

Usage :  python3 scripts/importer-classeur.py chemin/vers/classeur.xlsx
Besoin :  pip install openpyxl

Rien n'est inventé : tout champ vient d'une cellule du classeur. Une cellule
« — » devient null.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

SORTIE = Path(__file__).resolve().parent.parent / "data"

AGENTS = {
    "Acquisition & CA": "ACQUISITION_CA",
    "Finance & Rentabilité": "FINANCE_RENTABILITE",
    "RH & Organisation": "RH_ORGANISATION",
    "Comm & Création": "COMM_CREATION",
}
ONGLETS_DONNEES = {
    "Acquisition — données": "ACQUISITION_CA",
    "Finance — données": "FINANCE_RENTABILITE",
    "RH — données": "RH_ORGANISATION",
    "Comm — données": "COMM_CREATION",
}
NATURES = {"Marché": "marche", "Pratique": "pratique", "Cohérence": "coherence", "Contexte": "contexte"}
MODES = {"Demandée": "demandee", "Importée": "importee", "Calculée": "calculee", "Partagée": "partagee", "Autre agent": "autre_agent"}
NIVEAUX = {"Essentiel": "essentiel", "Approfondi": "approfondi"}
CHIFFRE = {"Oui": "oui", "Non": "non", "Calculé": "calcule"}
SUR = {"CA": "ca", "trésorerie": "tresorerie", "contacts": "contacts", "marge": "marge", "heures": "heures"}
# 8 profils depuis le classeur du 20 octobre (« Location » : hébergement et location de véhicules).
PROFILS = ["rdv", "commerce", "ecommerce", "restaurant", "chantier", "b2b", "abonnement", "location"]
# Les onglets Croisements n'ont pas encore de colonne Location : aucun croisement ne s'applique à ce profil.
PROFILS_CROISEMENTS = PROFILS[:7]
# Lexiques propres à un métier (onglet Lexique, colonnes après « Clé dans les questions »), clé technique.
LEXIQUES_METIER = {
    "Hébergement": "hebergement",
    "Location de véhicules": "location_vehicules",
    "Agence immobilière": "immobilier",
    "Projets de particuliers": "projet_particuliers",
    "Vente de biens": "vente_de_biens",
    "Négoce": "negoce",
    "Fabrication": "fabrication",
    "Cabinet": "cabinet",
    "Course (VTC, taxi, coursier)": "course",
    "Livraison de repas": "livraison_repas",
}
TAILLES = ["seul", "equipe_2_5", "structure_6_plus"]
RE_DONNEE = re.compile(r"\b(?:ACQ|FIN|RH|COM)-\d{2}\.\d{2}\b")
RE_CROISEMENT = re.compile(r"\b(?:XA|XF|XC|XR)\d{2}\b")
RE_DECLENCHEUR = re.compile(r"\bD\d{2}\b")


def v(cellule):
    """Texte d'une cellule, ou None si vide ou « — »."""
    if cellule is None:
        return None
    t = str(cellule).strip()
    return None if t in ("", "—") else t


def codes(texte):
    if texte is None:
        return []
    vus = []
    for c in RE_DONNEE.findall(texte) + RE_CROISEMENT.findall(texte):
        if c not in vus:
            vus.append(c)
    return vus


def lignes(ws, depuis=2):
    for r in ws.iter_rows(min_row=depuis, values_only=True):
        if r and r[0] not in (None, ""):
            yield r


def ecrire(nom, contenu):
    SORTIE.mkdir(exist_ok=True)
    (SORTIE / nom).write_text(json.dumps(contenu, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"  {nom} : {len(contenu)} éléments")


def themes(wb):
    sortie = []
    for r in lignes(wb["Thèmes"]):
        impact = v(r[5])
        if impact is None:
            impact_obj = {"type": "aucun", "sur": None}
        elif impact == "Qualifié":
            impact_obj = {"type": "qualifie", "sur": None}
        else:
            m = re.fullmatch(r"Chiffré \((.+)\)", impact)
            impact_obj = {"type": "chiffre", "sur": SUR[m.group(1)]}
        sortie.append({
            "code": r[0],
            "agent": AGENTS[r[1]],
            "bloc": r[2],
            "nom": r[3],
            "nature": NATURES[r[4]],
            "impact": impact_obj,
            "chiffreNecessaire": CHIFFRE[r[6]],
            "dejaConnu": v(r[8]),
            "methodeVolume": v(r[9]),
            "siInconnu": v(r[10]),
            "masquePour": v(r[11]),
        })
    ecrire("themes.json", sortie)


def metiers_concernes(texte, code):
    """Colonne « Métiers concernés » (module métier) : libellés de la liste des métiers séparés par « ; », une activité
    s'écrit « Métier / Activité ». Renvoie les identifiants (« vtc-et-taxi », « agence-immobiliere/transaction »)."""
    if texte is None:
        return None
    routage = json.loads((SORTIE / "routage.json").read_text(encoding="utf-8"))
    par_libelle = {m["libelle"].casefold(): m for m in routage["metiers"]}
    sortie = []
    for morceau in [x.strip() for x in str(texte).split(";") if x.strip()]:
        metier, _, activite = morceau.partition(" / ")
        m = par_libelle.get(metier.strip().casefold())
        if m is None:
            sys.exit(f"{code} : métier inconnu dans « Métiers concernés » : « {metier} ».")
        if not activite:
            sortie.append(m["id"])
            continue
        a = next((x for x in m.get("sousVariantes") or [] if x["libelle"].casefold() == activite.strip().casefold()), None)
        if a is None:
            sys.exit(f"{code} : activité inconnue pour {m['libelle']} : « {activite} ».")
        sortie.append(f'{m["id"]}/{a["id"]}')
    return sortie or None


def donnees(wb):
    for onglet, agent in ONGLETS_DONNEES.items():
        ws = wb[onglet]
        entetes = [c.value for c in ws[1]]
        rh = agent == "RH_ORGANISATION"
        cles = TAILLES if rh else PROFILS
        i0 = entetes.index("Seul") if rh else entetes.index("RDV")
        i_sert = entetes.index("Sert à")
        i_metiers = entetes.index("Métiers concernés") if "Métiers concernés" in entetes else None
        sortie = []
        for r in lignes(ws):
            marques = {k: (r[i0 + n] == "●") for n, k in enumerate(cles)}
            metiers = metiers_concernes(v(r[i_metiers]) if i_metiers is not None and i_metiers < len(r) else None, r[0])
            sortie.append({
                "code": r[0],
                "theme": r[1],
                "ecran": v(r[3]),
                "libelle": r[4],
                "format": v(r[5]),
                "mode": MODES[r[6]],
                "niveau": NIVEAUX.get(v(r[7]) or ""),
                "condition": v(r[8]),
                "declencheurs": RE_DECLENCHEUR.findall(r[8] or ""),
                "periode": v(r[9]),
                ("tailles" if rh else "profils"): marques,
                "sertA": [s.strip() for s in (v(r[i_sert]) or "").split(",") if s.strip()],
                "calcul": v(r[i_sert + 1]),
                "entrees": codes(v(r[i_sert + 2])),
                "entreesTexte": v(r[i_sert + 2]),
                "utiliseeDans": codes(v(r[i_sert + 3])),
                "utiliseeDansTexte": v(r[i_sert + 3]),
                **({"metiers": metiers} if metiers else {}),
            })
        ecrire(f"donnees-{agent.lower().replace('_', '-')}.json", sortie)


def croisements(wb):
    sortie = []
    for onglet in ("Croisements", "Croisements RH"):
        ws = wb[onglet]
        entetes = [c.value for c in ws[1]]
        rh = onglet == "Croisements RH"
        cles = TAILLES if rh else PROFILS_CROISEMENTS
        i0 = entetes.index("Seul") if rh else entetes.index("RDV")
        for r in lignes(ws):
            marques = {k: (r[i0 + n] == "●") for n, k in enumerate(cles)}
            if not rh:
                marques["location"] = False
            sortie.append({
                "code": r[0],
                "agent": AGENTS[r[1]],
                "nom": r[2],
                "calcul": r[3],
                "entrees": codes(v(r[4])),
                "entreesTexte": v(r[4]),
                ("tailles" if rh else "profils"): marques,
                "hypotheseExterne": v(r[i0 + len(cles)]),
                "phraseRapport": v(r[i0 + len(cles) + 1]),
            })
    ecrire("croisements.json", sortie)


def scenarios(wb):
    sortie = []
    for r in lignes(wb["Scénarios"], depuis=5):
        sortie.append({
            "code": r[0],
            "nom": r[1],
            "calcul": r[2],
            "entrees": codes(v(r[3])),
            "entreesTexte": v(r[3]),
            "phraseRapport": v(r[4]),
        })
    ecrire("scenarios.json", sortie)


def diagnostic_comm(wb):
    problemes = []
    ordre = []
    regle_vanite = None
    for r in lignes(wb["Diagnostic Comm"], depuis=5):
        if re.fullmatch(r"DC\d{2}", str(r[0])):
            problemes.append({
                "code": r[0],
                "probleme": r[1],
                "signaux": r[2],
                "entrees": codes(v(r[3])),
                "entreesTexte": v(r[3]),
                "recommande": r[4],
                "neRecommandePas": r[5],
            })
        elif str(r[0]).startswith("1."):
            ordre = [re.sub(r"^\d+\.\s*", "", p).strip() for p in re.split(r"\s{2}·\s{2}", str(r[0]))]
        elif str(r[0]).startswith("Abonnés"):
            regle_vanite = str(r[0])
    (SORTIE / "diagnostic-comm.json").write_text(
        json.dumps({"problemes": problemes, "ordreLectureChiffres": ordre, "regleChiffresDeVanite": regle_vanite}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"  diagnostic-comm.json : {len(problemes)} problèmes, ordre de lecture en {len(ordre)} niveaux")


def declencheurs(wb):
    sortie = []
    for r in lignes(wb["Déclencheurs"], depuis=5):
        sortie.append({
            "code": r[0],
            "question": r[1],
            "posee": r[2],
            "reponses": r[3],
            "effet": r[4],
            "donneesConcernees": int(r[5]) if r[5] not in (None, "") else None,
        })
    ecrire("declencheurs.json", sortie)


def cle_lexique(mot):
    """Même règle que `cleLexique` (src/lib/donnees/lexique.ts)."""
    t = unicodedata.normalize("NFD", mot.lower())
    t = re.sub(r"[\u0300-\u036f]", "", t)
    t = re.sub(r"['’]", " ", t)
    return re.sub(r"[^a-z0-9]+", "_", t).strip("_")


def lexique(wb):
    """Onglet Lexique : une ligne par mot ; colonnes = 7 profils, « Clé dans les questions », puis les lexiques par
    métier. Le profil « location » prend les formes de la colonne « Hébergement » ; les lexiques par métier sont
    gardés tels quels (champ `metiers`), le moteur les branchera métier par métier."""
    ws = wb["Lexique"]
    toutes = [r for r in ws.iter_rows(values_only=True)]
    i_entete = next(i for i, r in enumerate(toutes) if r[0] == "Mot générique")
    entete = [str(c).strip() if c is not None else "" for c in toutes[i_entete]]
    i_cle = entete.index("Clé dans les questions")
    i_metiers = {LEXIQUES_METIER[t]: i for i, t in enumerate(entete) if t in LEXIQUES_METIER}
    mots = []
    for r in toutes[i_entete + 1:]:
        if not r[0] or str(r[0]).startswith("Quel profil"):
            break
        formes = dict(zip(PROFILS[:7], [v(c) for c in r[1:8]]))
        metiers = {k: v(r[i]) for k, i in i_metiers.items()}
        formes["location"] = metiers.get("hebergement")
        mot = {"mot": r[0]}
        cle = (v(r[i_cle]) or "").strip("{} ")
        if cle and cle != cle_lexique(r[0]):
            mot["cle"] = cle
        mot["formes"] = formes
        mot["metiers"] = metiers
        mots.append(mot)
    ecrire("lexique.json", mots)
    return entete[1:8]


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    print("Import :", sys.argv[1])
    themes(wb)
    donnees(wb)
    croisements(wb)
    scenarios(wb)
    diagnostic_comm(wb)
    declencheurs(wb)
    lexique(wb)


if __name__ == "__main__":
    main()

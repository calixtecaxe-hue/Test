#!/usr/bin/env python3
"""Génère le gabarit de rédaction de la banque de questions : un classeur Excel avec une
ligne par donnée posée au dirigeant, déjà préremplie avec ce que le classeur « Thèmes »
permet de déduire (type de réponse, unité, options des tranches). Il reste à écrire la
question, l'aide, et à préciser les réponses que le format ne définit pas (Choix sans
liste, plusieurs champs, tableaux).

Usage :  python3 scripts/exporter-modele-banque.py [--sortie CAXE_Banque_Questions.xlsx] [--avec-questions]

Avec --avec-questions, les questions déjà écrites (data/banque/questions.json) sont reportées dans
le classeur, avec leur aide, leur repli et leur réponse attendue ; celles mises de côté
(data/banque/en-attente.json) sont ajoutées en fin de feuille, marquées « en attente ».
Ensuite : remplir le classeur, puis  python3 scripts/importer-banque.py CAXE_Banque_Questions.xlsx

Besoin :  pip install openpyxl  et  node (pour la déduction des réponses, une seule logique).
"""
import argparse
import json
import subprocess
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from banque_commun import (  # noqa: E402
    COLONNES, TYPES, TYPES_INVERSE, UNITES, UNITES_INVERSE,
    champs_vers_texte, options_vers_texte,
)

RACINE = Path(__file__).resolve().parent.parent
AGENTS = {
    "ACQUISITION_CA": "Acquisition",
    "FINANCE_RENTABILITE": "Finance",
    "RH_ORGANISATION": "RH",
    "COMM_CREATION": "Comm",
}


def charger():
    sortie = subprocess.run(["node", str(RACINE / "scripts" / "proposer-specs.mts")], capture_output=True, text=True, cwd=RACINE)
    if sortie.returncode != 0:
        sys.exit("Impossible de lire les données : " + sortie.stderr.strip()[:400])
    return json.loads(sortie.stdout)


def cellules_depuis_spec(spec):
    """Ce qu'on préremplit dans les colonnes de réponse pour une spécification sûre."""
    c = {"type": TYPES_INVERSE[spec["type"]]}
    t = spec["type"]
    if t == "nombre":
        c["unite"] = UNITES_INVERSE[spec["unite"]]
        if "min" in spec:
            c["min"] = spec["min"]
        if "max" in spec:
            c["max"] = spec["max"]
    elif t == "texte":
        if "max" in spec:
            c["max"] = spec["max"]
    elif t in ("choix", "cases"):
        c["options"] = options_vers_texte(spec["options"])
        if t == "cases":
            c["min"] = spec.get("minChoix")
            c["max"] = spec.get("maxChoix")
    elif t == "liste":
        c["max"] = spec["max"]
        c["min"] = spec.get("min")
    elif t == "composite":
        c["champs"] = champs_vers_texte(spec["champs"])
    elif t == "tableau":
        c["champs"] = champs_vers_texte(spec["colonnes"])
        lg = spec["lignes"]
        if lg["source"] == "fixe":
            c["lignes"] = "fixe: " + " / ".join(i["libelle"] for i in lg["items"])
        elif lg["source"] == "declencheur":
            c["lignes"] = f"cadrage: {lg['cle']}"
        elif lg["source"] == "donnee":
            c["lignes"] = f"liste: {lg['code']}"
        else:
            c["lignes"] = f"libre: {lg['max']}"
        if "total" in spec:
            ids = spec["total"].get("colonnes") or [spec["total"]["colonne"]]
            noms = [next(x for x in spec["colonnes"] if x["id"] == i)["libelle"] for i in ids]
            c["total"] = f"{' ; '.join(noms)} = {spec['total']['valeur']}"
    return c


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sortie", default=str(RACINE / "CAXE_Banque_Questions.xlsx"))
    ap.add_argument("--avec-questions", action="store_true")
    args = ap.parse_args()
    d = charger()
    ecrites = {}
    en_attente = []
    if args.avec_questions:
        ecrites = {q["donnee"]: q for q in json.loads((RACINE / "data" / "banque" / "questions.json").read_text(encoding="utf-8"))}
        en_attente = json.loads((RACINE / "data" / "banque" / "en-attente.json").read_text(encoding="utf-8"))

    wb = openpyxl.Workbook()
    # --- Mode d'emploi ---
    ws = wb.active
    ws.title = "Mode d'emploi"
    texte = [
        ("Banque de questions CAXE — mode d'emploi", True),
        ("", False),
        ("Une ligne = une donnée que le dirigeant renseigne. Écrivez la question dans la colonne « QUESTION » ; rien d'autre n'est obligatoire quand la réponse est déjà préremplie.", False),
        ("", False),
        ("Les colonnes grises viennent du classeur « Thèmes » : ne les modifiez pas. Les colonnes jaunes sont à vous.", False),
        ("", False),
        ("1. QUESTION. Une phrase, sans question rhétorique, sans réponse glissée dedans, en langage simple. Les mots qui changent selon le métier s'écrivent entre accolades : {demande}, {nouveau_client}, {devis}… (voir l'onglet Lexique). Une question conditionnelle se pose en deux temps : deux lignes.", False),
        ("2. Type de réponse. Déjà rempli quand le format du classeur le dit (Nombre, €, %, Tranches…). À remplir quand il est vide :", False),
        ("      Nombre : un chiffre. Choisissez l'unité. Min et Max sont facultatifs (ex. % entre 0 et 100).", False),
        ("      Oui / non  ·  Texte (Max = nombre de caractères au plus, facultatif)  ·  Mois et année", False),
        ("      Choix unique : une seule option parmi une liste. Écrivez les options séparées par « · » (ex. Hors zone · Budget trop faible · Hors offre).", False),
        ("      Cases à cocher : plusieurs options. Min et Max = nombre de cases à cocher.", False),
        ("      Liste libre : le dirigeant écrit lui-même une liste (ses types de clients). Max = nombre d'éléments au plus.", False),
        ("      Plusieurs champs : plusieurs valeurs d'une même question. Colonne « Champs » : Libellé [type] · Libellé [type]. Types : nombre, €, %, jours, heures, semaines, mois, ans, km, € / heure, oui / non, texte, mois et année, choix: A / B / C. Bornes : [% 0-100]  [€ 0-].", False),
        ("      Tableau : des lignes et des colonnes. Colonne « Champs » = les colonnes (même écriture). Colonne « Lignes » :", False),
        ("            fixe: Moins de 5 km / 5 à 10 km / 10 à 25 km     (lignes écrites d'avance)", False),
        ("            cadrage: D03     (une ligne par option cochée à la question de cadrage D03, D20 ou D22 : canaux, jours, réseaux)", False),
        ("            liste: ACQ-01.01     (une ligne par élément de la liste écrite à cette autre question, qui doit être de type Liste libre)", False),
        ("            libre: 5     (le dirigeant ajoute ses lignes, 5 au plus)", False),
        ("         « Total imposé » : par exemple  Part du CA = 100  oblige la colonne à faire 100.", False),
        ("3. « Je ne sais pas » et « Je ne suis pas cette donnée » sont proposés automatiquement sur toutes les questions : inutile de les écrire. Si le classeur liste « je ne sais pas » parmi des tranches, il est déjà retiré des options.", False),
        ("4. Aide : une phrase qui explique où trouver le chiffre. Si « je ne sais pas » : une phrase qui dit quoi faire pour le retrouver (ex. « Comptez les demandes pendant 4 semaines »).", False),
        ("", False),
        ("Une fois rempli : python3 scripts/importer-banque.py CAXE_Banque_Questions.xlsx — l'import contrôle tout, signale chaque erreur avec son numéro de ligne, et n'écrit rien tant qu'il en reste.", False),
        ("Vous pouvez remplir par morceaux : une question sans texte est simplement ignorée et reste signalée « manquante ».", False),
    ]
    for i, (t, gras) in enumerate(texte, 1):
        ws.cell(i, 1, t).font = Font(bold=gras, size=14 if gras else 11)
        ws.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 150

    # --- Questions ---
    wq = wb.create_sheet("Questions")
    gris = PatternFill("solid", fgColor="EDEDED")
    jaune = PatternFill("solid", fgColor="FFF8DC")
    for j, (_, titre, largeur, a_remplir) in enumerate(COLONNES, 1):
        c = wq.cell(1, j, titre)
        c.font = Font(bold=True)
        c.fill = jaune if a_remplir else gris
        c.alignment = Alignment(wrap_text=True, vertical="center")
        wq.column_dimensions[get_column_letter(j)].width = largeur
    wq.row_dimensions[1].height = 48
    wq.freeze_panes = "B2"

    cle_col = {cle: j for j, (cle, *_r) in enumerate(COLONNES, 1)}
    ordre = {"ACQUISITION_CA": 0, "FINANCE_RENTABILITE": 1, "RH_ORGANISATION": 2, "COMM_CREATION": 3}
    lignes = sorted(d["lignes"], key=lambda x: (ordre[x["agent"]], x["code"]))
    for i, l in enumerate(lignes, 2):
        valeurs = {
            "code": l["code"],
            "agent": AGENTS[l["agent"]],
            "theme": l["theme"],
            "niveau": "Essentiel" if l["niveau"] == "essentiel" else "Approfondi",
            "format": l["format"],
            "libelle": l["libelle"],
            "concerne": ", ".join(l["concerne"]),
            "condition": l["condition"],
            "indice": l["proposition"]["indice"],
        }
        spec = l["proposition"]["spec"]
        if spec:
            valeurs.update(cellules_depuis_spec(spec))
        q = ecrites.get(l["code"])
        if q:
            valeurs.update({"question": q["texte"], "aide": q.get("aide"), "repli": q.get("repli")})
            for cle in ("type", "unite", "min", "max", "options", "champs", "lignes", "total"):
                valeurs.pop(cle, None)
            valeurs.update(cellules_depuis_spec(q["reponse"]))
        for cle, v in valeurs.items():
            c = wq.cell(i, cle_col[cle], v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if not next(x for x in COLONNES if x[0] == cle)[3]:
                c.fill = gris
        for cle, _t, _w, a_remplir in COLONNES:
            if a_remplir:
                wq.cell(i, cle_col[cle]).alignment = Alignment(wrap_text=True, vertical="top")
                if wq.cell(i, cle_col[cle]).value is None:
                    wq.cell(i, cle_col[cle]).fill = jaune
    for k, q in enumerate(en_attente, len(lignes) + 2):
        valeurs = {
            "code": q["donnee"],
            "agent": {"ACQ": "Acquisition", "FIN": "Finance", "RH-": "RH", "COM": "Comm"}[q["donnee"][:3]],
            "libelle": "EN ATTENTE : cette donnée n'est pas encore dans le classeur des 4 agents (" + (q.get("raison") or "à ajouter") + ")",
            "question": q["texte"], "aide": q.get("aide"), "repli": q.get("repli"),
        }
        valeurs.update(cellules_depuis_spec(q["reponse"]))
        for cle, v in valeurs.items():
            c = wq.cell(k, cle_col[cle], v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    fin = len(lignes) + len(en_attente) + 1

    dv_type = DataValidation(type="list", formula1='"' + ",".join(TYPES.keys()) + '"', allow_blank=True)
    dv_unite = DataValidation(type="list", formula1='"' + ",".join(UNITES.keys()) + '"', allow_blank=True)
    wq.add_data_validation(dv_type)
    wq.add_data_validation(dv_unite)
    dv_type.add(f"{get_column_letter(cle_col['type'])}2:{get_column_letter(cle_col['type'])}{fin}")
    dv_unite.add(f"{get_column_letter(cle_col['unite'])}2:{get_column_letter(cle_col['unite'])}{fin}")
    wq.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES))}{fin}"

    # --- Lexique ---
    wl = wb.create_sheet("Lexique")
    profils = ["rdv", "commerce", "ecommerce", "restaurant", "chantier", "b2b", "abonnement"]
    for j, t in enumerate(["Mot à écrire dans la question", "Mot"] + profils, 1):
        wl.cell(1, j, t).font = Font(bold=True)
    import unicodedata, re
    for i, m in enumerate(d["lexique"], 2):
        s = m.get("cle") or re.sub(r"[^a-z0-9]+", "_", "".join(c for c in unicodedata.normalize("NFD", m["mot"].lower()) if not unicodedata.combining(c))).strip("_")
        wl.cell(i, 1, "{" + s + "}")
        wl.cell(i, 2, m["mot"])
        for j, p in enumerate(profils, 3):
            wl.cell(i, j, m["formes"].get(p) or "—")
    wl.column_dimensions["A"].width = 34
    wl.column_dimensions["B"].width = 28
    for j in range(3, 10):
        wl.column_dimensions[get_column_letter(j)].width = 22

    # --- Cadrage ---
    wc = wb.create_sheet("Cadrage")
    for j, t in enumerate(["Code", "Question de cadrage (écrite dans le classeur)", "Options (lignes d'un tableau « cadrage: Dxx »)"], 1):
        wc.cell(1, j, t).font = Font(bold=True)
    for i, dc in enumerate(d["declencheurs"], 2):
        wc.cell(i, 1, dc["code"])
        wc.cell(i, 2, dc["question"])
        wc.cell(i, 3, " · ".join(dc["options"]))
    wc.column_dimensions["A"].width = 8
    wc.column_dimensions["B"].width = 70
    wc.column_dimensions["C"].width = 90

    wb.save(args.sortie)
    prerempli = sum(1 for l in lignes if l["proposition"]["spec"])
    print(f"{args.sortie} : {len(lignes)} questions à écrire, dont {prerempli} dont le type de réponse est déjà rempli et {len(lignes) - prerempli} à préciser.")


if __name__ == "__main__":
    main()
